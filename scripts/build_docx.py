#!/usr/bin/env python3
"""out/INTX_report.md -> out/INTX_report.docx

의존적인 마크다운 파서를 쓰지 않고, 보고서가 실제로 사용하는 문법만 처리한다:
  # ~ ##### 제목 / 문단 / - · 1. 목록 / GFM 표 / ``` 코드·도식 블록 /
  > 인용 / --- 구분선 / **굵게** `코드` [R###]

[R###] 태그는 INTX_evidence.html 의 해당 앵커로 하이퍼링크된다.
"""

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "out", "INTX_report.md")
OUT = os.path.join(ROOT, "out", "INTX_report.docx")
EVIDENCE_HTML = "INTX_evidence.html"

KO_FONT = "맑은 고딕"
MONO = "D2Coding"

RID_RE = re.compile(r"\[(R\d{3,4}(?:\s*,\s*R\d{3,4})*)\]")
EST_RE = re.compile(r"\[EST\]")
TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[R\d{3,4}(?:\s*,\s*R\d{3,4})*\]|\[EST\])")


def set_ko_font(run, name=KO_FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), name)


def add_hyperlink(paragraph, url, text, color="1D4ED8", size=None, bold=False):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")

    rf = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), MONO)
    rpr.append(rf)

    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    if bold:
        rpr.append(OxmlElement("w:b"))
    r.append(rpr)

    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def emit_inline(par, text, size=10.5, base_bold=False, color=None):
    """**굵게**, `코드`, [R###] 링크, [EST] 배지를 처리해 run 으로 쪼갠다."""
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            run = par.add_run(tok[2:-2])
            run.bold = True
            run.font.size = Pt(size)
            set_ko_font(run)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            run = par.add_run(tok[1:-1])
            run.font.size = Pt(size - 0.5)
            set_ko_font(run, MONO)
            run.font.color.rgb = RGBColor.from_string("A32020")
        elif tok == "[EST]":
            run = par.add_run("[EST]")
            run.bold = True
            run.font.size = Pt(size - 1)
            set_ko_font(run, MONO)
            run.font.color.rgb = RGBColor.from_string("8A6100")
        elif RID_RE.fullmatch(tok):
            inner = tok[1:-1]
            par.add_run("[").font.size = Pt(size - 1)
            rids = [x.strip() for x in inner.split(",")]
            for i, rid in enumerate(rids):
                if i:
                    par.add_run(",").font.size = Pt(size - 1)
                add_hyperlink(par, f"{EVIDENCE_HTML}#{rid}", rid, size=size - 1)
            par.add_run("]").font.size = Pt(size - 1)
        else:
            run = par.add_run(tok)
            run.bold = base_bold
            run.font.size = Pt(size)
            set_ko_font(run)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:fill"), hexcolor)
    tcpr.append(sh)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_sep(line):
    return bool(re.fullmatch(r"\s*\|?[\s:\-\|]+\|?\s*", line)) and "-" in line


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    t = doc.add_table(rows=len(rows), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(1)
            par.paragraph_format.space_after = Pt(1)
            emit_inline(par, cell_text.replace("<br>", " / "), size=8.5, base_bold=(i == 0))
            if i == 0:
                shade(cell, "EDF0F3")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def add_pre(doc, lines):
    """도식/코드 블록: 단일 셀 표에 고정폭으로 넣어 정렬을 보존."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade(cell, "F6F7F9")
    cell.text = ""
    first = True
    for ln in lines:
        par = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        par.paragraph_format.space_before = Pt(0)
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        run = par.add_run(ln if ln.strip() else " ")
        run.font.size = Pt(7.5)
        set_ko_font(run, MONO)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main():
    if not os.path.exists(SRC):
        sys.exit(f"[FATAL] 보고서 원본이 없습니다: {SRC}")

    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().split("\n")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.size = Pt(10.5)
    style.font.name = KO_FONT
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), KO_FONT)

    # 페이지 하단 쪽번호
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 코드/도식 블록
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            if buf:
                add_pre(doc, buf)
            continue

        # 표
        if stripped.startswith("|") and i + 1 < n and is_sep(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # 구분선
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            p = doc.add_paragraph()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "DFE3E8")
            pbdr.append(bottom)
            p._p.get_or_add_pPr().append(pbdr)
            i += 1
            continue

        # 제목
        m = re.match(r"^(#{1,5})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1:
                doc.add_page_break() if len(doc.paragraphs) > 3 else None
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt({1: 14, 2: 14, 3: 10, 4: 8, 5: 6}[level])
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            size = {1: 18, 2: 14.5, 3: 12.5, 4: 11, 5: 10.5}[level]
            color = {1: "0F1216", 2: "1D4ED8", 3: "16191D", 4: "5C6570", 5: "5C6570"}[level]
            emit_inline(p, text, size=size, base_bold=True, color=color)
            if level <= 2:
                pbdr = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6" if level == 1 else "4")
                bottom.set(qn("w:color"), "DFE3E8")
                pbdr.append(bottom)
                p._p.get_or_add_pPr().append(pbdr)
            i += 1
            continue

        # 인용
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(4)
            emit_inline(p, stripped.lstrip("> ").strip(), size=10, color="5C6570")
            i += 1
            continue

        # 목록
        m = re.match(r"^(\s*)([-*•]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            ordered = bool(re.match(r"\d+\.", m.group(2)))
            p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            p.paragraph_format.left_indent = Cm(0.6 + 0.5 * indent)
            p.paragraph_format.space_after = Pt(2)
            emit_inline(p, m.group(3).strip())
            i += 1
            continue

        # 문단
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.35
        emit_inline(p, stripped)
        i += 1

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc.save(OUT)

    src_words = sum(len(l.split()) for l in lines)
    print(f"[OK] {OUT}")
    print(f"     원본 {len(lines)}줄 / 약 {src_words} 단어")
    print(f"     [R###] 태그는 {EVIDENCE_HTML} 앵커로 하이퍼링크됨 (같은 폴더에 두고 열 것)")


if __name__ == "__main__":
    main()
